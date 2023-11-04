
import yaml
import openpyxl
from openpyxl.chart import PieChart, BarChart, Reference

def ReadConfigs(yaml_path):
    data_dict = {}
    with open(yaml_path, encoding='UTF8') as f:
        data_dict = yaml.load(f, Loader=yaml.FullLoader)
    return data_dict

class ComputeTime():
    def __init__(self, data_dict):
        self.answer_list = []
        self.ave_list = data_dict["Average"]
        self.question = data_dict["Question_list"]
        self.maximum_life = int(data_dict["Maximum_Life"])
        self.kid_maximun_age = int(data_dict["Kid_Maximum_Age"])
        self.remaining_time = 0
        self.gender = ""
        self.shortage_in_year = 0
        self.shortage_in_hour = 0   
        self.out_name = data_dict["Out_Name"] 
        
        pass
    
    def __compute_year_to_hour__(self):
        self.shortage_in_hour = self.shortage_in_year * 365 * 24
        
    def __ask_gender__(self):
        print("자녀와의 관계를 입력해주세요.")
        answer = input("아버지 / 어머니: ")
        if answer in ["아버지", "어머니"]:
            self.gender = answer
        else:
            print("아버지 또는 어머니로 입력해주세요.")
            return self.__ask_gender__()
        
    def __question__(self, idx):
        if (idx >= len(self.question)):
            return 0
        try:
            print(self.question[idx], end="")
            answer = input(": ")
            answer = int(answer)
            return answer
        except:
            print("정수를 입력하세요.")
            return self.__question__(idx)
        
    def __compute_q1__(self):
        age_myself = self.__question__(0)
        age_myself_shortage = self.maximum_life - age_myself
        age_kid = self.__question__(1)
        age_kid_shortage = self.kid_maximun_age - age_kid
        if (age_kid_shortage < 0):
            print("이 설문은 {} 세 이하의 자녀를 가진 부모를 대상으로 작성되었습니다.".format(self.kid_maximun_age))
            exit()
        self.shortage_in_year = min(age_kid_shortage, age_myself_shortage)
        self.__compute_year_to_hour__()
        pass
    
    def __compute_q2__(self):
        hour_sleep = self.__question__(2)
        self.answer_list.append(hour_sleep)
        self.shortage_in_hour = self.shortage_in_hour - hour_sleep * 365 * self.shortage_in_year
        pass
    def __compute_q3__(self):
        min_moving = self.__question__(3)
        self.answer_list.append(min_moving)
        self.shortage_in_hour = self.shortage_in_hour - min_moving * 1.0 / 60.0 * 365 * self.shortage_in_year
        pass
    def __compute_q4__(self):
        hour_working = self.__question__(4)
        self.answer_list.append(hour_working)
        self.shortage_in_hour = self.shortage_in_hour - hour_working * 365 * self.shortage_in_year
        pass
    def __compute_q5__(self):
        hour_relax = self.__question__(5)
        self.answer_list.append(hour_relax)
        self.shortage_in_hour = self.shortage_in_hour - hour_relax * 365 * self.shortage_in_year
        pass
    def __compute_q6__(self):
        hour_with_kid = self.__question__(6)
        self.answer_list.append(hour_with_kid)
        pass
    
    def __time_with_kid__(self, hour_with_kid):
        data = {
            0.5: 15.7,
            2: 43.6,
            24: 40.6
        }
        for k, v in data.items():
            if hour_with_kid < k:
                return v
    def __generate_report__(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["내 수면시간", "평균 수면시간"])
        sheet.append([self.answer_list[0], self.ave_list[0]])
        
        sheet.append(["내 이동시간", "평균 이동시간"])
        sheet.append([self.answer_list[1], self.ave_list[1]])
        
        sheet.append(["내 노동시간", "평균 노동시간"])
        sheet.append([self.answer_list[2], self.ave_list[2]])
        
        sheet.append(["내 여가시간", "평균 여가시간"])
        sheet.append([self.answer_list[3], self.ave_list[3]])
        
        left_hour = 24
        left_hour -= self.answer_list[0]
        left_hour -= self.answer_list[1] / 60.
        left_hour -= self.answer_list[2]
        left_hour -= self.answer_list[3]
        data = [
            ['항목', '시간'],
            ['수면시간', self.answer_list[0]],
            ['이동시간', self.answer_list[1] / 60.],
            ['노동시간', self.answer_list[2]],
            ['여가시간', self.answer_list[3]],
            ['육아시간', self.answer_list[4]],
            ['남은시간', left_hour]
        ]
        
        for row in data:
            sheet.append(row)
            
        values = Reference(sheet, min_col=1, min_row=1
                           ,max_col=2, max_row=2)
        chart = BarChart()
        chart.add_data(values, titles_from_data=True)
        chart.title = "평균 수면시간 대비 내 수면시간"
        chart.y_axis.title = "수면시간"
        chart.x_axis.title = "단위: 시간"
        sheet.add_chart(chart, "A1")
        
        
        values = Reference(sheet, min_col=1, min_row=3
                           ,max_col=2, max_row=4)
        chart = BarChart()
        chart.add_data(values, titles_from_data=True)
        chart.title = "평균 이동시간 대비 내 이동시간"
        chart.y_axis.title = "이동시간"
        chart.x_axis.title = "단위: 분"
        sheet.add_chart(chart, "I1")
        
        
        values = Reference(sheet, min_col=1, min_row=5
                           ,max_col=2, max_row=6)
        chart = BarChart()
        chart.add_data(values, titles_from_data=True)
        chart.title = "평균 노동시간 대비 내 노동시간"
        chart.y_axis.title = "노동시간"
        chart.x_axis.title = "단위: 시간"
        sheet.add_chart(chart, "A14")
        
        
        values = Reference(sheet, min_col=1, min_row=7
                           ,max_col=2, max_row=8)
        chart = BarChart()
        chart.add_data(values, titles_from_data=True)
        chart.title = "평균 여가시간 대비 내 여가시간"
        chart.y_axis.title = "여가시간"
        chart.x_axis.title = "단위: 시간"
        sheet.add_chart(chart, "I14")
        
        pie = PieChart()
        labels = Reference(sheet, min_col = 1, min_row = 10, max_row=15)
        data = Reference(sheet, min_col = 2, min_row = 9, max_row=15)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = "하루 시간표"
        sheet.add_chart(pie, "A27")
        
        percents = self.__time_with_kid__(self.answer_list[4])
        word = "당신의 자녀와 하루 놀이/대화시간은 전체 부모의 {}%에 속합니다.".format(percents)
        sheet['J28'] = word
        wb.save(self.out_name)
        pass 
        
        
        
    def __call__(self):
        self.__ask_gender__()
        self.__compute_q1__()
        self.__compute_q2__()
        self.__compute_q3__()
        self.__compute_q4__()
        self.__compute_q5__()
        self.__compute_q6__()
        print("당신이 자녀와 함께 할 수 있는 시간은 총 {} 시간 남았습니다.".format(int(self.shortage_in_hour)))
        print("당신이 자녀와 함께 할 수 있는 시간은 총 {} 일 남았습니다.".format(int(self.shortage_in_hour / 24)))
        print("당신이 자녀와 함께 할 수 있는 시간은 총 {} 년 남았습니다.".format(int(self.shortage_in_hour / (24 * 365))))
        self.__generate_report__()
        pass
        
    
if __name__=="__main__":
    yaml_path = "Configure.yaml"
    computeTime = ComputeTime(ReadConfigs(yaml_path))
    computeTime()
    